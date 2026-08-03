from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CLIENT_ROOT = Path(__file__).resolve().parents[3]
OUT = CLIENT_ROOT / "01-Source-Material" / "access-guides" / "Hetzner_VPS_Read_Only_Access_Guide.docx"

NAVY = "17365D"
BLUE = "2E74B5"
INK = "1F2937"
MUTED = "5B6573"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREEN = "EAF6EF"
GREEN = "23633A"
LIGHT_GOLD = "FFF7DF"
GOLD = "7A5A00"
LIGHT_GRAY = "F4F6F8"
BORDER = "D7DEE7"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para(p, before=0, after=6, line=1.25, keep_next=False):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_next


def add_text(doc, text, *, bold_prefix=None, after=6, color=INK, italic=False):
    p = doc.add_paragraph()
    set_para(p, after=after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, bold=True, color=color)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, color=color, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_callout(doc, label, text, *, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=accent, size="8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.15)
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True, color=accent)
    r = p.add_run(text)
    set_font(r, color=INK)
    spacer = doc.add_paragraph()
    set_para(spacer, after=2, line=1.0)
    return table


def add_code(doc, lines):
    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        set_para(p, before=0, after=0 if idx < len(lines) - 1 else 6, line=1.05)
        p.paragraph_format.keep_with_next = idx < len(lines) - 1
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_GRAY)
        p_pr.append(shd)
        r = p.add_run(line if line else " ")
        set_font(r, name="Courier New", size=8.4, color="20242A")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_numbering(doc, num_fmt, text, left=540, hanging=270, color=INK):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_fmt)
    lvl.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_item(doc, text, num_id, *, bold_prefix=None, after=4, keep_next=False):
    p = doc.add_paragraph()
    set_para(p, after=after, line=1.25)
    p.paragraph_format.keep_with_next = keep_next
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_field(paragraph, instruction):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instruction)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.36)
section.footer_distance = Inches(0.38)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, NAVY, 8, 4),
):
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_para(hp, after=0, line=1.0)
run = hp.add_run("ACCESS GUIDE")
set_font(run, size=8.5, color=BLUE, bold=True)
run = hp.add_run("   |   OPENCLAW PLATFORM REVIEW")
set_font(run, size=8.5, color=MUTED)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_para(fp, after=0, line=1.0)
run = fp.add_run("Darsh Ahirrao  |  Temporary access guide  |  Page ")
set_font(run, size=8.3, color=MUTED)
add_field(fp, "PAGE")
run = fp.add_run(" of ")
set_font(run, size=8.3, color=MUTED)
add_field(fp, "NUMPAGES")

kicker = doc.add_paragraph()
set_para(kicker, before=8, after=2, line=1.0)
run = kicker.add_run("CLIENT ACCESS INSTRUCTIONS")
set_font(run, size=9, color=BLUE, bold=True)

title = doc.add_paragraph()
set_para(title, before=0, after=5, line=1.0, keep_next=True)
run = title.add_run("Temporary Read-Only Access\nfor the Hetzner VPS")
set_font(run, size=27, color=NAVY, bold=True)

subtitle = doc.add_paragraph()
set_para(subtitle, after=14, line=1.15)
run = subtitle.add_run("The easiest safe setup for the initial OpenClaw platform technical review")
set_font(run, size=12.5, color=MUTED)

meta = doc.add_table(rows=2, cols=2)
set_table_geometry(meta, [4680, 4680])
set_table_borders(meta, color=BORDER, size="5")
metadata = [
    (0, 0, "Prepared for", "Mark Gerhart"),
    (0, 1, "Prepared by", "Darsh Ahirrao"),
    (1, 0, "Access type", "Temporary, key-only, no sudo"),
    (1, 1, "Use this guide", "After the Upwork contract is active"),
]
for row, col, label, value in metadata:
    cell = meta.cell(row, col)
    set_cell_shading(cell, WHITE if row == 0 else LIGHT_GRAY)
    p = cell.paragraphs[0]
    set_para(p, after=0, line=1.05)
    r = p.add_run(label.upper() + "\n")
    set_font(r, size=8, color=MUTED, bold=True)
    r = p.add_run(value)
    set_font(r, size=10.2, color=INK, bold=True)

spacer = doc.add_paragraph()
set_para(spacer, after=2, line=1.0)

add_callout(
    doc,
    "Recommended setup",
    "Create one temporary SSH user for Darsh, use his public key, give the account no sudo access, and grant read/traverse permission only to the two agreed content-platform directories.",
    fill=LIGHT_GREEN,
    accent=GREEN,
)

add_heading(doc, "Before you begin", 1)
bullet_num = add_numbering(doc, "bullet", "•", left=540, hanging=270)
for text in (
    "Keep your current administrator SSH session open until the new login has been tested.",
    "Ask Darsh for one public-key line beginning with ssh-ed25519. A public key is safe to install; never ask him for a private key.",
    "Confirm the exact absolute paths for Sapphire and Sable before granting folder access.",
    "Check the selected folders for .env files, private keys, tokens, credentials, or unrelated client data. Exclude or redact secrets before granting recursive read access.",
):
    add_list_item(doc, text, bullet_num)

add_callout(
    doc,
    "Important",
    "For an existing Hetzner VPS, the key must be installed inside the server account. Adding a key only in the Hetzner Console does not retroactively add it to an existing server.",
    fill=LIGHT_GOLD,
    accent=GOLD,
)

doc.add_page_break()
setup_heading = add_heading(doc, "Setup steps", 1)
step_num = add_numbering(doc, "decimal", "%1.", left=540, hanging=270)

add_list_item(doc, "Create a dedicated temporary user named darsh-review.", step_num, after=3, keep_next=True)
add_code(doc, [
    'sudo adduser --disabled-password --gecos "" darsh-review',
    "sudo passwd -l darsh-review",
])

add_list_item(doc, "Create the SSH key folder with strict permissions.", step_num, after=3, keep_next=True)
add_code(doc, [
    "sudo install -d -m 700 -o darsh-review -g darsh-review /home/darsh-review/.ssh",
    "sudo nano /home/darsh-review/.ssh/authorized_keys",
])
add_text(doc, "Paste the single public-key line supplied by Darsh. In nano, press Ctrl+O, Enter, then Ctrl+X.", after=4, color=MUTED)
add_code(doc, [
    "sudo chown darsh-review:darsh-review /home/darsh-review/.ssh/authorized_keys",
    "sudo chmod 600 /home/darsh-review/.ssh/authorized_keys",
])

add_list_item(doc, "Set the two approved project paths.", step_num, after=3, keep_next=True)
add_code(doc, [
    'SAPPHIRE_DIR="/replace/with/exact/sapphire/path"',
    'SABLE_DIR="/replace/with/exact/sable/path"',
    'printf "%s\\n%s\\n" "$SAPPHIRE_DIR" "$SABLE_DIR"',
])
add_text(doc, "Replace both placeholders before continuing. If either path is uncertain, stop and contact Darsh.", after=4, color=MUTED)

add_list_item(doc, "Install the ACL utility if needed, then grant read/traverse access only.", step_num, after=3, keep_next=True)
add_code(doc, [
    "command -v setfacl >/dev/null || sudo apt-get install -y acl",
    'sudo setfacl -R -m u:darsh-review:r-X "$SAPPHIRE_DIR" "$SABLE_DIR"',
])
add_text(doc, "The r-X permission reads files and traverses directories without granting project write permission. Do not add default ACLs unless Darsh specifically requests them.", after=4, color=MUTED)

add_list_item(doc, "Verify that the account can read but cannot write.", step_num, after=3, keep_next=True)
add_code(doc, [
    'sudo -u darsh-review ls -la "$SAPPHIRE_DIR" >/dev/null && echo "Sapphire: readable"',
    'sudo -u darsh-review ls -la "$SABLE_DIR" >/dev/null && echo "Sable: readable"',
    'sudo -u darsh-review find "$SAPPHIRE_DIR" "$SABLE_DIR" -writable -print -quit',
    "sudo -l -U darsh-review",
])
add_text(doc, "The find command should print nothing. The sudo check should show that darsh-review is not allowed to run sudo.", after=4, color=MUTED)

add_list_item(doc, "Send Darsh the connection details through Upwork after the contract starts.", step_num, after=3, keep_next=True)
send_num = add_numbering(doc, "bullet", "•", left=720, hanging=270)
for text in (
    "Server IP address or hostname",
    "SSH username: darsh-review",
    "SSH port, if it is not 22",
    "The two approved project paths",
    "Confirmation that no sudo access was granted",
):
    add_list_item(doc, text, send_num, after=2)

doc.add_page_break()
permissions_heading = add_heading(doc, "What this account can and cannot do", 1)
add_text(doc, "It can inspect the files and directories explicitly granted above and can write only inside its own temporary home directory.", bold_prefix="It can ")
add_text(doc, "It cannot modify the approved project folders, restart services, install packages, change system configuration, or use sudo.", bold_prefix="It cannot ")
add_text(doc, "Some service-manager output and protected logs may remain unavailable. If Darsh needs one of those items, he will request the smallest specific read-only export or permission required.")

add_callout(
    doc,
    "Do not broaden access",
    "Do not add darsh-review to sudo, root, adm, systemd-journal, Docker, or any application-owner group. Do not grant access to the Crypto Intelligence platform unless it is separately approved in scope.",
    fill=LIGHT_GOLD,
    accent=GOLD,
)

add_heading(doc, "Remove access when the review is complete", 1)
add_text(doc, "Run these commands from an administrator account after Darsh confirms the review is complete:")
add_code(doc, [
    'sudo setfacl -R -x u:darsh-review "$SAPPHIRE_DIR" "$SABLE_DIR"',
    "sudo userdel -r darsh-review",
])
add_text(doc, "If additional parent-directory ACLs were added during setup, remove those entries as well.", after=8, color=MUTED)

add_callout(
    doc,
    "Need help?",
    "If any command, folder path, permission, or service layout is unclear, stop and contact Darsh. Do not guess or expand access. Darsh can walk you through the setup and confirm the minimum permissions needed.",
    fill=LIGHT_GREEN,
    accent=GREEN,
)

add_heading(doc, "Official references", 2)
refs = [
    ("Ubuntu: OpenSSH server and public-key authentication", "https://ubuntu.com/server/docs/how-to/security/openssh-server/"),
    ("Ubuntu: setfacl access-control-list manual", "https://manpages.ubuntu.com/manpages/jammy/man1/setfacl.1.html"),
    ("Hetzner: SSH keys on existing servers", "https://docs.hetzner.com/cloud/servers/how-to-rescue/change-ssh-key/"),
]
for label, url in refs:
    p = doc.add_paragraph()
    set_para(p, after=2, line=1.1)
    add_hyperlink(p, label, url)

doc.core_properties.title = "Temporary Read-Only Access for the Hetzner VPS"
doc.core_properties.subject = "Client instructions for scoped SSH access"
doc.core_properties.author = "Darsh Ahirrao"
doc.core_properties.keywords = "Hetzner, SSH, read-only access, OpenClaw, technical review"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
