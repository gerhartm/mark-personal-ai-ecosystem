from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CLIENT_ROOT = Path(__file__).resolve().parents[3]
WORKSPACE = Path(__file__).resolve().parents[2]
OUTPUT_DOCX = CLIENT_ROOT / "04-Deliverables" / "option-1"
OUTPUT_PDF = CLIENT_ROOT / "04-Deliverables" / "option-1"
WORK = WORKSPACE / ".work" / "proposal"
DOCX_PATH = OUTPUT_DOCX / "Mark-Gerhart-Option-1-Architecture-and-Phase-Plan.docx"
DIAGRAM_PATH = WORK / "personal-ai-ecosystem-architecture.png"

OUTPUT_DOCX.mkdir(parents=True, exist_ok=True)
OUTPUT_PDF.mkdir(parents=True, exist_ok=True)
WORK.mkdir(parents=True, exist_ok=True)


# Design tokens: standard_business_brief + customer_pack header pattern.
NAVY = "102A43"
BLUE = "176B87"
TEAL = "168A8A"
INK = "1F2933"
MUTED = "627D98"
LIGHT = "F2F6F8"
LIGHT_BLUE = "E8F2F6"
LIGHT_TEAL = "E7F5F3"
LIGHT_GRAY = "F4F6F8"
BORDER = "C7D3DA"
FUTURE = "8AA1AE"
WHITE = "FFFFFF"
GOLD = "B7791F"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size=6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_in: list[float], indent_dxa=120) -> None:
    assert abs(sum(widths_in) - 6.5) < 0.001
    widths_dxa = [round(width * 1440) for width in widths_in]
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_in[idx])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_run_font(run, *, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def add_bottom_border(paragraph, color=BLUE, size=12) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_paragraph_shading(paragraph, fill: str, left_border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if left_border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def no_split(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_para(
    doc,
    text="",
    *,
    style=None,
    size=None,
    bold=False,
    italic=False,
    color=INK,
    align=None,
    before=0,
    after=6,
    line=1.1,
):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size or 11, color=color, bold=bold, italic=italic)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    left = 0.48 if level == 0 else 0.72
    p.paragraph_format.left_indent = Inches(left)
    p.paragraph_format.first_line_indent = Inches(-0.24)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    marker = p.add_run("•  ")
    set_run_font(marker, size=9.8, color=INK)
    r = p.add_run(text)
    set_run_font(r, size=9.8)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    return p


def add_callout(doc, label: str, text: str, fill=LIGHT_BLUE, border=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.12
    set_paragraph_shading(p, fill, border)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=NAVY, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    # Style carries format; direct formatting ensures renderer consistency.
    sizes = {1: 16, 2: 13, 3: 12}
    colors = {1: BLUE, 2: BLUE, 3: NAVY}
    set_run_font(r, size=sizes[level], color=colors[level], bold=True)
    no_split(p)
    return p


def add_metadata_table(doc):
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [3.25, 3.25])
    items = [
        ("Client", "Mark Gerhart"),
        ("Engagement", "Option 1 - Crypto Intelligence Second Brain"),
        ("Status", "Technical architecture and implementation plan"),
    ]
    for row, (label, value) in zip(table.rows, items):
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label.upper())
        set_run_font(r0, size=8.5, color=MUTED, bold=True)
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(value)
        set_run_font(r1, size=9.5, color=INK, bold=True)
    return table


def rounded_box(draw, xy, *, fill, outline, width=4, radius=24):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def font(size, bold=False):
    path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf"
    return ImageFont.truetype(path, size)


def wrap(draw, text, max_width, font_obj):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font_obj)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def center_text(draw, box, title, details=(), *, title_color="#102A43", body_color="#334E68"):
    x1, y1, x2, y2 = box
    title_font = font(31, True)
    body_font = font(24, False)
    title_lines = wrap(draw, title, x2 - x1 - 40, title_font)
    body_lines = []
    for detail in details:
        body_lines.extend(wrap(draw, detail, x2 - x1 - 40, body_font))
    total_h = len(title_lines) * 39 + (10 if body_lines else 0) + len(body_lines) * 31
    y = y1 + (y2 - y1 - total_h) / 2
    for line in title_lines:
        bbox = draw.textbbox((0, 0), line, font=title_font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=title_font, fill=title_color)
        y += 39
    y += 10
    for line in body_lines:
        bbox = draw.textbbox((0, 0), line, font=body_font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=body_font, fill=body_color)
        y += 31


def arrow(draw, start, end, *, color="#627D98", width=5, dashed=False):
    x1, y1 = start
    x2, y2 = end
    if dashed:
        steps = max(1, int(math.hypot(x2 - x1, y2 - y1) / 22))
        for i in range(steps):
            if i % 2 == 0:
                t1 = i / steps
                t2 = min(1, (i + 1) / steps)
                draw.line((x1 + (x2 - x1) * t1, y1 + (y2 - y1) * t1,
                           x1 + (x2 - x1) * t2, y1 + (y2 - y1) * t2),
                          fill=color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 18
    wing = 0.55
    p1 = (x2 - length * math.cos(angle - wing), y2 - length * math.sin(angle - wing))
    p2 = (x2 - length * math.cos(angle + wing), y2 - length * math.sin(angle + wing))
    draw.polygon([(x2, y2), p1, p2], fill=color)


def build_architecture_diagram():
    image = Image.new("RGB", (2100, 1660), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    # Top access flow
    channels = (90, 90, 560, 270)
    access = (760, 90, 1340, 270)
    outputs = (1540, 90, 2010, 270)
    rounded_box(draw, channels, fill="#F2F6F8", outline="#8AA1AE")
    rounded_box(draw, access, fill="#E8F2F6", outline="#176B87")
    rounded_box(draw, outputs, fill="#F2F6F8", outline="#8AA1AE")
    center_text(draw, channels, "CHANNELS", ("Telegram", "Web workspace", "Uploads and future channels"))
    center_text(draw, access, "SECURE APPLICATION LAYER", ("Cloudflare protection", "Authentication and API"))
    center_text(draw, outputs, "USER OUTCOMES", ("Cited answers", "Timelines", "Briefs and content"))
    arrow(draw, (560, 180), (760, 180), color="#176B87")

    # Central brain
    brain = (530, 370, 1570, 610)
    rounded_box(draw, brain, fill="#DDF2F1", outline="#168A8A", width=6, radius=34)
    center_text(
        draw,
        brain,
        "ONE HERMES CENTRAL BRAIN",
        ("Reasoning, planning, workspace routing, tool selection and human approval rules",),
        title_color="#0E5F5F",
    )
    arrow(draw, (1050, 270), (1050, 370), color="#168A8A")
    arrow(draw, (1570, 490), (1780, 270), color="#176B87")

    # Project branches
    branch_y1, branch_y2 = 720, 950
    boxes = [
        ((70, branch_y1, 520, branch_y2), "#D9EEF5", "#176B87", "CRYPTO INTELLIGENCE", ("BUILD NOW", "Research, sourced answers", "timelines and content")),
        ((585, branch_y1, 1035, branch_y2), "#F4F6F8", "#8AA1AE", "AI TOOLING", ("FUTURE BRANCH", "Tool capture, comparison", "and recommendations")),
        ((1100, branch_y1, 1550, branch_y2), "#F4F6F8", "#8AA1AE", "REAL ESTATE CONTENT", ("FUTURE BRANCH", "Examples, hooks, scripts", "and content planning")),
        ((1615, branch_y1, 2065, branch_y2), "#F4F6F8", "#8AA1AE", "FUTURE PROJECT N", ("EXTENSION POINT", "New skills, workflows", "and permissions")),
    ]
    for idx, (box, fill, outline, title, details) in enumerate(boxes):
        rounded_box(draw, box, fill=fill, outline=outline, width=5)
        center_text(draw, box, title, details, title_color=outline)
        arrow(draw, (1050, 610), ((box[0] + box[2]) / 2, box[1]), color=outline, dashed=idx != 0)

    # Ultimate memory
    memory = (360, 1065, 1740, 1325)
    rounded_box(draw, memory, fill="#FFF7E6", outline="#B7791F", width=6, radius=34)
    center_text(
        draw,
        memory,
        "ONE ULTIMATE MEMORY",
        (
            "Global identity and preferences | Project branches | Source provenance",
            "Facts, entities, events and relationships | Approved feedback | Skills and decisions",
        ),
        title_color="#8A5A14",
    )
    for idx, (box, _, outline, _, _) in enumerate(boxes):
        start = ((box[0] + box[2]) / 2, box[3])
        end = (510 + idx * 360, memory[1])
        arrow(draw, start, end, color=outline, dashed=idx != 0)

    # Storage and operations band
    stores = [
        ((80, 1430, 500, 1585), "POSTGRESQL", ("Canonical records", "permissions and history")),
        ((565, 1430, 985, 1585), "PGVECTOR", ("Semantic retrieval", "derived search index")),
        ((1050, 1430, 1470, 1585), "OBJECT STORAGE", ("Documents, media", "and transcripts")),
        ((1535, 1430, 2020, 1585), "OPERATIONS", ("Coolify, workers, monitoring", "and off-server backups")),
    ]
    for box, title, details in stores:
        rounded_box(draw, box, fill="#F2F6F8", outline="#627D98", width=4, radius=20)
        center_text(draw, box, title, details)
    for box, _, _ in stores[:3]:
        arrow(draw, (1050, memory[3]), ((box[0] + box[2]) / 2, box[1]), color="#B7791F")
    arrow(draw, (1740, memory[3] - 30), (1775, 1430), color="#627D98")

    # Active/future legend
    active_font = font(22, True)
    draw.rounded_rectangle((80, 1610, 205, 1640), radius=10, fill="#D9EEF5", outline="#176B87", width=3)
    draw.text((225, 1607), "Solid path: implemented in Option 1", font=active_font, fill="#334E68")
    draw.line((905, 1625, 1035, 1625), fill="#8AA1AE", width=4)
    for i in range(6):
        if i % 2 == 0:
            draw.line((905 + i * 22, 1625, 925 + i * 22, 1625), fill="#8AA1AE", width=6)
    draw.text((1060, 1607), "Dashed path: designed extension, not implemented now", font=active_font, fill="#334E68")

    image.save(DIAGRAM_PATH, quality=96, dpi=(180, 180))


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name, left, first in (
        ("List Bullet", 0.50, -0.25),
        ("List Bullet 2", 0.75, -0.25),
        ("List Number", 0.50, -0.25),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(left)
        style.paragraph_format.first_line_indent = Inches(first)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("PERSONAL AI ECOSYSTEM  |  OPTION 1 IMPLEMENTATION PLAN")
    set_run_font(r, size=8.2, color=MUTED, bold=True)
    add_bottom_border(hp, color=BORDER, size=4)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [4.8, 1.7], indent_dxa=0)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.first_child_found_in("w:tcBorders")
        if borders is not None:
            tc_pr.remove(borders)
    p0 = table.rows[0].cells[0].paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run("Confidential | Prepared for Mark Gerhart")
    set_run_font(r0, size=8, color=MUTED)
    p1 = table.rows[0].cells[1].paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run("Page ")
    set_run_font(r1, size=8, color=MUTED)
    add_field(p1, "PAGE")


def add_cover(doc: Document):
    add_para(doc, "IMPLEMENTATION PROPOSAL", size=10, bold=True, color=TEAL, after=14)
    p = add_para(doc, "Personal AI Ecosystem", size=30, bold=True, color=NAVY, after=4)
    p.paragraph_format.keep_with_next = True
    add_para(doc, "Option 1: Crypto Intelligence Second Brain", size=17, bold=True, color=BLUE, after=14)
    add_para(
        doc,
        "One Hermes central brain. One ultimate memory. One complete Crypto Intelligence capability built first on an expandable foundation.",
        size=12.5,
        color=MUTED,
        after=18,
        line=1.18,
    )
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(18)
    add_bottom_border(rule, color=TEAL, size=16)
    add_metadata_table(doc)
    add_para(doc, "", after=7)
    add_callout(
        doc,
        "Decision in this document",
        "Approve the technical implementation phases and completion criteria for Option 1. The plan is capped at 100 combined implementation hours and preserves the complete existing VPS before any cutover.",
        fill=LIGHT_TEAL,
        border=TEAL,
    )
    add_para(doc, "25 July 2026", size=9.5, color=MUTED, italic=True, after=2)
    doc.add_page_break()


def add_executive_summary(doc: Document):
    add_heading(doc, "1. Proposed Direction", 1)
    add_para(
        doc,
        "Option 1 will establish the shared foundation for Mark's Personal AI Ecosystem and deliver one complete, production-ready Crypto Intelligence second brain. Hermes will operate as the single central brain. A single Ultimate Memory service will preserve global context, project knowledge, source provenance, feedback and reusable procedures.",
    )
    add_para(
        doc,
        "Crypto Intelligence is the only project branch implemented during Option 1. AI Tooling, Real Estate Content and future projects are represented as designed extension points so they can be added later without rebuilding the central brain, memory model, security, deployment or ingestion foundation.",
    )
    add_callout(
        doc,
        "Outcome",
        "Mark can submit research through Telegram or the web workspace, retrieve source-grounded answers, generate event timelines, and create panel or social content from accumulated knowledge.",
    )

    add_heading(doc, "Guiding Principles", 2)
    for item in (
        "One Hermes brain rather than separate agents with drifting identities.",
        "One logical memory with scoped project branches rather than disconnected databases.",
        "Every retained fact remains traceable to its original source.",
        "Automation handles routine ingestion, organization and retrieval; consequential actions remain reviewable.",
        "Existing Crypto, Sable and Sapphire content is preserved before the current VPS is changed.",
        "Deployment, health, backups and resource limits are designed into the first release.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Option 1 Boundary", 2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [3.25, 3.25])
    set_repeat_table_header(table.rows[0])
    headers = ("Implemented now", "Designed for later")
    for idx, text in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=WHITE, bold=True)
    row = table.add_row()
    left = [
        "Shared brain and memory foundation",
        "Crypto ingestion and knowledge workflows",
        "Crypto web and Telegram experience",
        "Crypto data migration, launch and recovery",
    ]
    right = [
        "AI Tooling-specific workflows and interface",
        "Real Estate-specific workflows and interface",
        "Sable and Sapphire data activation",
        "Additional project workspaces",
    ]
    for idx, values in enumerate((left, right)):
        p = row.cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for j, value in enumerate(values):
            if j:
                p = row.cells[idx].add_paragraph()
            p.style = doc.styles["List Bullet"]
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(value)
            set_run_font(r, size=9.4)
    doc.add_page_break()


def add_architecture(doc: Document):
    add_heading(doc, "2. Target Architecture", 1)
    add_para(
        doc,
        "The ecosystem is organized as one central intelligence with multiple capability branches. Solid paths in the diagram are delivered in Option 1. Dashed paths are intentional extension points and are not included in the current implementation scope.",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(5)
    inline_picture = p.add_run().add_picture(str(DIAGRAM_PATH), width=Inches(6.35))
    inline_picture._inline.docPr.set(
        "descr",
        "Personal AI Ecosystem architecture: channels connect through a secure application layer to one Hermes central brain. Crypto Intelligence is implemented first, while AI Tooling, Real Estate Content and future projects remain extension branches. All branches use one Ultimate Memory backed by PostgreSQL, pgvector, object storage and Coolify-managed operations.",
    )
    inline_picture._inline.docPr.set("title", "Personal AI Ecosystem target architecture")
    caption = add_para(
        doc,
        "Figure 1. Complete Personal AI Ecosystem architecture with Crypto Intelligence activated first.",
        size=9,
        italic=True,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    no_split(caption)
    doc.add_page_break()


def add_architecture_components(doc: Document):
    add_heading(doc, "3. Architecture Components", 1)
    add_heading(doc, "Single Hermes Central Brain", 2)
    add_para(
        doc,
        "Hermes provides one continuous identity for reasoning, planning, tool selection and workspace routing. Telegram, the web workspace and future channels reach the same central brain. Project behavior is supplied by scoped skills, tools and memory context rather than by creating separate autonomous brains.",
    )
    add_callout(
        doc,
        "Concurrency rule",
        "Option 1 uses one active Hermes gateway. Media extraction, transcription, indexing and scheduled processing scale through separate workers so parallel workloads do not compete with the central brain or share its local state unsafely.",
        fill=LIGHT_GRAY,
        border=MUTED,
    )

    add_heading(doc, "Ultimate Memory", 2)
    add_para(
        doc,
        "Ultimate Memory is one logical source of truth exposed through a controlled Memory API. It is not one oversized prompt or a single unstructured file. Hermes receives a compact context package containing global identity, the active project branch and task-relevant memories.",
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [1.55, 2.35, 2.60])
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(("Memory layer", "Purpose", "Implementation")):
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9, color=WHITE, bold=True)
    rows = [
        ("Global memory", "Identity, preferences, goals and approved operating rules", "Small curated profile supplied to every relevant session"),
        ("Project branches", "Crypto now; AI Tooling, Real Estate and future projects later", "Scoped records, skills, permissions and retrieval filters"),
        ("Source vault", "Original documents, URLs, media and transcripts", "Object storage with immutable identifiers and provenance"),
        ("Knowledge graph", "Facts, entities, events, relationships and timelines", "PostgreSQL canonical records with version history"),
        ("Retrieval index", "Fast semantic and keyword recall", "pgvector index derived from canonical memory"),
        ("Memory operations", "Capture, deduplicate, consolidate, correct and export", "Controlled write service, audit log and Markdown views"),
    ]
    for values in rows:
        row = table.add_row()
        for idx, value in enumerate(values):
            if len(table.rows) % 2 == 1:
                set_cell_shading(row.cells[idx], LIGHT_GRAY)
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=8.7, color=INK, bold=idx == 0)

    add_heading(doc, "Shared Processing and Tool Layer", 2)
    for item in (
        "Workspace-aware ingestion assigns every source to a project branch before processing.",
        "A background queue controls concurrency, retries and resource-heavy work.",
        "Adapters acquire supported URLs and media; workers extract, transcribe, normalize and index content.",
        "A controlled tool gateway exposes retrieval, timelines and generation without giving the agent unrestricted database access.",
        "Each memory write includes project scope, source, timestamp, confidence, permissions and version history.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Deployment and Operations", 2)
    add_para(
        doc,
        "Coolify acts as the control plane for deployments, domains, environment configuration and service health. The application stack remains reproducible from version-controlled Docker Compose definitions. Application data is backed up independently from the Coolify control plane.",
    )
    doc.add_page_break()


def add_existing_vps(doc: Document):
    add_heading(doc, "4. Existing VPS Preservation and Migration", 1)
    add_para(
        doc,
        "The current VPS contains application data and behavior from Crypto Intelligence, Sable and Sapphire. Option 1 does not reset or overwrite that server at the beginning of the engagement. The entire environment is preserved first so it remains a reference, migration source and rollback point.",
    )
    add_heading(doc, "Preserve Everything, Migrate Selectively", 2)
    steps = [
        ("Snapshot", "Create and verify a provider-level snapshot plus an encrypted off-server backup."),
        ("Inventory", "Record databases, uploaded media, transcripts, events, content, conversation archives, sidecar state and relevant configuration."),
        ("Classify", "Mark each item as migrate now, preserve for later, or exclude as temporary/system-generated material."),
        ("Import Crypto", "Transform relevant Crypto content into the new canonical memory and object-storage model."),
        ("Reconcile", "Compare source and target counts, preserve timestamps and source identifiers, and verify representative records."),
        ("Hold for rollback", "Keep the original VPS read-only until Mark accepts the new Crypto system and migration."),
    ]
    for idx, (title, body) in enumerate(steps, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Inches(0.08)
        r1 = p.add_run(f"{idx}. {title}. ")
        set_run_font(r1, size=10.5, color=BLUE, bold=True)
        r2 = p.add_run(body)
        set_run_font(r2, size=10.5, color=INK)

    add_callout(
        doc,
        "Option 1 migration rule",
        "Crypto content is migrated and activated now. Sable and Sapphire content remains fully preserved in the read-only archive for activation during a later ecosystem expansion.",
        fill=LIGHT_TEAL,
        border=TEAL,
    )

    add_heading(doc, "Migration Acceptance Checks", 2)
    for item in (
        "Original source files remain recoverable.",
        "Record counts are reconciled and exceptions are documented.",
        "Source URLs, timestamps and provenance survive migration.",
        "A representative sample opens, searches and answers correctly in V2.",
        "A database and object-storage restore test succeeds.",
        "No production cutover occurs before written acceptance.",
    ):
        add_bullet(doc, item)
    doc.add_page_break()


PHASES = [
    {
        "name": "Phase 1 - Infrastructure and Deployment Control Plane",
        "hours": "14-18",
        "goal": "Provision the isolated production environment and deploy the reproducible service foundation.",
        "work": [
            "Provision the fresh VPS with restricted SSH access, firewall rules and non-root administration.",
            "Deploy Coolify and connect the production stack to version-controlled Docker Compose definitions.",
            "Configure Cloudflare ingress, TLS, authentication, API routing, private service networks and secrets handling.",
            "Deploy PostgreSQL with pgvector, object storage, Redis-backed queues, health checks, monitoring and off-server backups.",
        ],
        "done": "The complete base stack deploys from Git, external traffic is authenticated and encrypted, service health is observable, and a backup can be restored without using the Coolify control plane.",
    },
    {
        "name": "Phase 2 - Hermes Runtime and Governed Memory Core",
        "hours": "14-20",
        "goal": "Deploy the single Hermes runtime and implement the controlled Ultimate Memory services.",
        "work": [
            "Deploy one active Hermes gateway with scoped tool execution, workspace routing and explicit permission boundaries.",
            "Implement the Memory API and canonical PostgreSQL schema for global context, project-scoped records, provenance and version history.",
            "Build hybrid retrieval using pgvector semantic search plus structured and keyword filters derived from canonical records.",
            "Implement context compilation, correction handling, audit logging and object-storage references for source files.",
        ],
        "done": "Hermes can store, retrieve and correct Crypto knowledge through the Memory API, return source provenance, preserve scoped context across sessions and deny unauthorized tool or workspace access.",
    },
    {
        "name": "Phase 3 - Multi-Source Ingestion and Processing Services",
        "hours": "14-20",
        "goal": "Convert submitted sources into durable, traceable and searchable knowledge through fault-tolerant workers.",
        "work": [
            "Implement authenticated Telegram and web ingestion adapters with request validation and project assignment.",
            "Process text, URLs, PDFs and agreed audio/video sources through bounded asynchronous workers.",
            "Normalize metadata, extract or transcribe content, deduplicate sources, chunk records and generate retrieval indexes.",
            "Implement idempotent job states, retry policies, dead-letter handling, progress visibility and actionable failure reporting.",
        ],
        "done": "Every agreed representative source type completes an end-to-end ingestion test, appears once in the library, remains linked to its original source and exposes visible processing or failure state.",
    },
    {
        "name": "Phase 4 - Crypto Intelligence Application and Workflows",
        "hours": "18-24",
        "goal": "Deliver the authenticated Crypto research workspace and production intelligence workflows.",
        "work": [
            "Build the authenticated conversational research interface, source library and ingestion-status views.",
            "Implement retrieval-augmented answers with citations, source inspection and project-scoped memory context.",
            "Implement event and entity extraction for chronological timeline generation from selected knowledge.",
            "Implement panel briefs, research summaries, X threads and LinkedIn content with reusable output controls.",
            "Write approved preferences, corrections and workflow feedback back into Ultimate Memory.",
        ],
        "done": "Mark can independently ingest sources, retrieve cited answers, inspect evidence, generate timelines and export each agreed content format from the production workspace.",
    },
    {
        "name": "Phase 5 - Data Migration, Hardening and Production Cutover",
        "hours": "10-18",
        "goal": "Migrate the agreed Crypto data, harden the complete stack and cut over with a tested rollback path.",
        "work": [
            "Snapshot and inventory the existing VPS, then export and transform the agreed Crypto records and source files.",
            "Import migrated data through controlled jobs and reconcile record counts, timestamps, identifiers and representative samples.",
            "Run functional, authorization, resource, failure-recovery, backup and restore acceptance tests.",
            "Configure production alerts and resource limits, resolve launch defects and document supported third-party fallbacks.",
            "Perform staged cutover, provide operating notes and preserve the old VPS through the agreed rollback window.",
        ],
        "done": "Migration is reconciled, all acceptance checks pass, recovery is tested, Mark accepts the production release and the old VPS can be retired after the agreed rollback window.",
    },
]


def add_phase_plan(doc: Document):
    add_heading(doc, "5. Technical Implementation Plan", 1)
    add_para(
        doc,
        "The architecture direction and initial VPS assessment are already complete, so they are excluded as implementation phases. The remaining work is organized into five technical phases covering deployment, core services, ingestion, application workflows, migration and production launch. Option 1 remains capped at 100 combined implementation hours.",
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table_geometry(table, [0.65, 2.55, 0.75, 2.55])
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(("Phase", "Outcome", "Hours", "Completion signal")):
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=8.5, color=WHITE, bold=True)
    summaries = [
        ("1", "Infrastructure and deployment control plane", "14-18", "Git-based production stack passes security, health and restore checks"),
        ("2", "Hermes runtime and governed memory core", "14-20", "Scoped memory, provenance and permission boundaries pass"),
        ("3", "Multi-source ingestion and processing", "14-20", "Representative sources complete traceable, idempotent processing"),
        ("4", "Crypto application and intelligence workflows", "18-24", "Research, citation, timeline and content workflows pass"),
        ("5", "Migration, hardening and production cutover", "10-18", "Data reconciled, recovery tested and production accepted"),
    ]
    for idx, values in enumerate(summaries):
        row = table.add_row()
        if idx % 2:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
        for col, value in enumerate(values):
            p = row.cells[col].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=8.4, color=INK, bold=col in (0, 2))
    total = table.add_row()
    for cell in total.cells:
        set_cell_shading(cell, LIGHT_TEAL)
    p = total.cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("TOTAL")
    set_run_font(r, size=9, color=NAVY, bold=True)
    p = total.cells[1].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Option 1 implementation")
    set_run_font(r, size=9, color=NAVY, bold=True)
    p = total.cells[2].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("70-100")
    set_run_font(r, size=9.5, color=NAVY, bold=True)
    p = total.cells[3].paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Hard cap: 100 hours")
    set_run_font(r, size=9, color=NAVY, bold=True)

    doc.add_page_break()
    add_heading(doc, "6. Technical Phase Completion Criteria", 1)
    for idx, phase in enumerate(PHASES):
        add_heading(doc, phase["name"], 2)
        p = add_para(doc, phase["goal"], size=10.5, color=MUTED, italic=True, after=5)
        no_split(p)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("Work included")
        set_run_font(r, size=9, color=NAVY, bold=True)
        for item in phase["work"]:
            add_bullet(doc, item)
        add_callout(doc, "Done when", phase["done"], fill=LIGHT_BLUE, border=BLUE)
        if idx in (1, 3):
            doc.add_page_break()


def add_scope_and_delivery(doc: Document):
    doc.add_page_break()
    add_heading(doc, "7. Scope, Assumptions and Delivery", 1)
    add_heading(doc, "Included in Option 1", 2)
    for item in (
        "Fresh VPS production environment, Coolify control plane and Git-based deployment stack.",
        "Single Hermes runtime and governed Ultimate Memory services required for Crypto.",
        "Complete Crypto ingestion, retrieval, timeline and content workflows.",
        "Web workspace and Telegram interaction for agreed users.",
        "Preservation of the complete existing VPS and migration of agreed Crypto content.",
        "Deployment definitions, monitoring, resource controls, backups, testing and handover.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Not Included in Option 1", 2)
    for item in (
        "Implementation of AI Tooling, Real Estate Content or other future project branches.",
        "Activation or migration of Sable and Sapphire into production V2 workspaces.",
        "Autonomous cryptocurrency trading, financial execution or investment decisions.",
        "Native mobile applications or integrations outside the agreed source and output list.",
        "Hosting, model-provider, storage, transcription or other third-party consumption charges.",
        "Guaranteed acquisition from third-party sites that actively block automation; supported fallbacks and clear failures are included.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Client Inputs Required", 2)
    for item in (
        "Confirmation of the delivery schedule and technical implementation plan.",
        "Approval of partner-assisted delivery and manual time for Darshan's agreed non-trackable work.",
        "Read-only access to the existing VPS until preservation and migration are complete.",
        "A clean deployment environment, domain/DNS access and approved provider accounts.",
        "Telegram bot access and confirmation of authorized users.",
        "Representative Crypto sources, evaluation questions and preferred output examples.",
        "A decision on which existing data must be migrated into Option 1.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Delivery Schedule Options", 2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [1.35, 2.25, 2.90])
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(("Schedule", "Combined weekly capacity", "Planning basis")):
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9, color=WHITE, bold=True)
    for idx, values in enumerate((
        ("10 days", "Up to 70 combined hours", "Partner-assisted fast-track delivery"),
        ("Two weeks", "Up to 50 combined hours", "Partner-assisted accelerated delivery"),
    )):
        row = table.add_row()
        if idx % 2:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_GRAY)
        for col, value in enumerate(values):
            p = row.cells[col].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=9.2, color=INK, bold=col == 0)

    add_callout(
        doc,
        "Delivery staffing",
        "Both schedules use combined capacity across Darshan and a development partner. Manual time must be enabled for agreed work that cannot be reliably captured through Time Tracker.",
        fill=LIGHT_BLUE,
        border=BLUE,
    )

    add_callout(
        doc,
        "Commercial guardrail",
        "Option 1 remains capped at 100 combined implementation hours under either schedule. Weekly capacity can be reduced after completion. Once the schedule, manual time and required access are confirmed, Phase 1 begins with the fresh VPS and deployment control plane. The existing VPS remains unchanged until its snapshot and migration inventory are verified.",
        fill=LIGHT_TEAL,
        border=TEAL,
    )


def build_docx():
    build_architecture_diagram()
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "Personal AI Ecosystem - Option 1 Architecture and Phase Plan"
    doc.core_properties.subject = "Crypto Intelligence Second Brain Implementation Proposal"
    doc.core_properties.author = "Darshan Ahirrao"
    doc.core_properties.keywords = "Hermes, Personal AI, Crypto Intelligence, architecture, implementation plan"

    add_cover(doc)
    add_executive_summary(doc)
    add_architecture(doc)
    add_architecture_components(doc)
    add_existing_vps(doc)
    add_phase_plan(doc)
    add_scope_and_delivery(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    print(DIAGRAM_PATH)


if __name__ == "__main__":
    build_docx()
